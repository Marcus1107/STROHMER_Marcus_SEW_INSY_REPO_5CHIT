from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor

out=Path(__file__).parent/'Textanalyse_SRDP.docx'
d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2)
s.left_margin=s.right_margin=Cm(2.2)
for name,size in [('Normal',11),('Title',23),('Heading 1',15),('Heading 2',12)]:
 st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,1)
d.add_paragraph('Textanalyse bei der SRDP','Title')
p('Lernunterlage für Deutsch mit Aufbau, Merkmalen und Textbausteinen')
p('Eine Textanalyse untersucht einen nicht fiktionalen Text, etwa einen Zeitungsartikel, einen Kommentar oder eine Rede. Sie erklärt anhand von Textbelegen, wie Inhalt, Aufbau und Sprache zusammenwirken. Diese Lernhilfe ist vollständig kopierbar und bearbeitbar.')
h('1 Wortanzahl')
p('405–495 oder 540–660 Wörter laut offiziellem Textsortenkatalog. Verbindlich ist die Wortspanne der konkreten Aufgabenstellung. Die angegebenen Grenzen enthalten bereits den Spielraum; keine zusätzlichen zehn Prozent aufschlagen.')
h('2 Aufbau schematisch')
p('Einleitung → knappe Inhaltsübersicht → Analyse von Aufbau, Argumentation und Sprache → zusammenfassendes Ergebnis')
p('Einleitung: Textsorte, Titel, Autor oder Autorin, Medium und Erscheinungsdatum nennen, soweit bekannt. Das Thema knapp angeben. Fehlende Angaben nicht erfinden.')
p('Hauptteil: Zunächst die für den Auftrag wichtigen Inhalte kurz darstellen. Anschließend untersuchen, wie der Text aufgebaut ist, wie die Argumentation funktioniert und welche sprachlichen Merkmale auffallen. Inhalt, Form und mögliche Wirkung miteinander verbinden. Die Arbeitsaufträge bestimmen die Schwerpunkte; nicht jeden denkbaren Aspekt abarbeiten.')
p('Schluss: Die wichtigsten Analyseergebnisse bündeln. Erklären, wie die untersuchten Mittel die Textfunktion unterstützen. Falls verlangt, die Wirksamkeit anhand der Ergebnisse einschätzen. Keine neuen Analysepunkte und kein persönliches Geschmacksurteil einführen.')
h('3 Inhalt und Argumentation')
p('Das 3-B-Schema lässt sich für die Analyse als Beobachtung beziehungsweise Behauptung → Beleg → Bedeutung oder Funktion nutzen. Das übliche argumentative Schema „Behauptung – Begründung – Beispiel“ wird hier also auf die Textuntersuchung angepasst.')
p('Fiktives Beispiel: „Der Verfasser formuliert seine Forderung nachdrücklich. Dies zeigt die Wiederholung von ‚Wir brauchen‘ am Beginn mehrerer Sätze. Die Anapher hebt die einzelnen Forderungen hervor und verbindet sie zu einer gemeinsamen Botschaft.“')
p('Pro und Kontra: Keine eigene Diskussion über das Sachthema schreiben. Stattdessen untersuchen, welche Position die Vorlage vertritt, welche Gründe sie nennt und ob sie Gegenargumente aufgreift, entkräftet oder ausblendet.')
p('Passend (+): zentrale These, Gedankengang, Argumente und Belege untersuchen; Beispiele und Expertenverweise auf ihre Funktion prüfen; Stilmittel korrekt benennen und in ihrem Zusammenhang erklären.')
p('Vermeiden (−): bloße Nacherzählung; unbelegte Vermutungen; eigene Pro-und-Kontra-Erörterung; Stilmittel ohne Erklärung aufzählen; pauschale Aussagen wie „Das regt zum Nachdenken an“.')

d.add_page_break()
h('4 Sprachliche Merkmale')
p('Die eigene Analyse: sachlich, präzise und überwiegend im Präsens schreiben. Fachbegriffe korrekt verwenden. Eigene Wertungen wie „Ich finde den Artikel toll“ vermeiden. Aussagen der Vorlage klar zuordnen; bei indirekter Rede kann der Konjunktiv I sinnvoll sein.')
p('Belege: Kurze direkte Zitate in Anführungszeichen setzen und mit Zeilenangaben belegen, wenn die Vorlage Zeilen vorgibt. Sinngemäße Verweise ebenfalls kenntlich machen. Zitate müssen genau stimmen; sie ersetzen die Erklärung nicht.')
p('In der Vorlage untersuchen: Wortwahl und Wortfelder, Fach- oder Umgangssprache, wertende Begriffe, Satzlänge und Satzbau, Fragen, direkte Ansprache, Wiederholungen und bildhafte Sprache. Nur auffällige, für den Auftrag relevante Merkmale auswählen.')
p('Mögliche Stilmittel und Funktionen:')
for t in [
'Anapher: gleicher Beginn aufeinanderfolgender Sätze oder Satzteile; kann zentrale Gedanken betonen.',
'Metapher: bildhafte Übertragung, etwa „eine Flut von Nachrichten“; kann eine große Menge anschaulich machen und zugleich bewerten.',
'Antithese: Gegenüberstellung gegensätzlicher Begriffe oder Gedanken; kann einen Konflikt zuspitzen.',
'Rhetorische Frage: Frage, auf die keine tatsächliche Antwort erwartet wird; kann Zustimmung nahelegen oder einen Einwand hervorheben.',
'Ellipse: unvollständiger, aus dem Zusammenhang verständlicher Satz; kann eine Aussage verdichten oder Dringlichkeit vermitteln.'
]: p(t)
p('Die Funktion hängt vom Kontext ab. Formuliere nachvollziehbar: „Die Wortwahl stellt … als bedrohlich dar“ oder „Die Wiederholung kann … hervorheben“. Behaupte keine bei allen Leserinnen und Lesern identische Wirkung.')
h('5 Abgrenzung')
p('Zur Zusammenfassung: Sie beantwortet vor allem „Was sagt der Text?“. Die Analyse untersucht zusätzlich „Wie ist der Text gestaltet und welche Funktion haben diese Merkmale?“.')
p('Zur Textinterpretation: Die Interpretation erschließt die Bedeutung eines literarischen Textes. Die Textanalyse bleibt bei nachvollziehbaren Befunden zu einem nicht fiktionalen Text und erklärt daraus dessen Gestaltung und Funktion.')
p('Zur Erörterung: Die Erörterung entwickelt eine eigene begründete Antwort auf eine Sachfrage. Die Textanalyse untersucht, wie die Vorlage ihre Aussagen und Argumente vermittelt.')
p('Zum Kommentar: Ein Kommentar vertritt eine eigene zugespitzte Position. In der Textanalyse ist die Position des untersuchten Textes Gegenstand der Untersuchung.')

d.add_page_break()
h('6 Textbausteine')
p('Die folgenden Sätze sind Formulierungshilfen. Angaben in eckigen Klammern ersetzen und nur Bausteine verwenden, die zur Vorlage und zum Arbeitsauftrag passen.')
blocks=[
('Einleitung',[
'Der [Textsorte] „[Titel]“ von [Name], erschienen am [Datum] in [Medium], behandelt [Thema].',
'Im Mittelpunkt steht die Frage, [Fragestellung].']),
('Inhalt und Aufbau',[
'Die zentrale Aussage des Textes lautet, dass [Kernaussage].',
'Zu Beginn beschreibt die Autorin [Aspekt]; anschließend geht sie auf [Aspekt] ein.',
'Der Übergang von [Abschnitt] zu [Abschnitt] markiert einen Wechsel von [Funktion] zu [Funktion].']),
('Argumentation',[
'Der Verfasser stützt seine These auf [Beleg oder Argument].',
'Der Verweis auf [Expertin oder Experte] soll die Aussage [Aussage] absichern.',
'Den Einwand, dass [Einwand], greift die Autorin auf und entgegnet, dass [Entgegnung].']),
('Sprache und Funktion',[
'Auffällig ist die Verwendung von [Merkmal], etwa in „[Zitat]“ (Z. [Zahl]).',
'Die Metapher „[Zitat]“ (Z. [Zahl]) veranschaulicht [Aspekt], indem sie [Erklärung].',
'Die wiederholte Formulierung „[Zitat]“ betont [Gedanke] und unterstützt damit [Aussage].',
'Mit der direkten Anrede „[Zitat]“ (Z. [Zahl]) bezieht der Text die Leserschaft in [Zusammenhang] ein.',
'Die wertenden Begriffe [Beispiele] stellen [Sachverhalt] als [Bewertung] dar.']),
('Schluss',[
'Zusammenfassend prägen vor allem [Merkmal] und [Merkmal] die Gestaltung des Textes.',
'Die untersuchten Mittel unterstützen die Absicht, [belegte Textfunktion], insbesondere durch [zentrales Ergebnis].'])]
for title,items in blocks:
 d.add_heading(title,2)
 for t in items: p('„'+t+'“')
d.add_heading('Quelle und Verwendung',2)
p('Wortspannen und Textsortenabgrenzung: offizieller Textsortenkatalog zur SRDP in der Unterrichtssprache, Fassung 2020, S. 16. Die Aufbauskizze und Beispielsätze sind eigenständig formulierte Lernhilfen. Maßgeblich bleibt der jeweilige Arbeitsauftrag.')
p('https://www.matura.gv.at/index.php?eID=dumpFile&f=4525&t=f&token=950c7f2b86f0ebc3459c5f0aa0e04013ab99c572')
d.core_properties.title='Textanalyse bei der SRDP'
d.save(out)
print(out.resolve())
